#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Флаги
  --limit N            показать только первые N совпадений
  --case-sensitive      учитывать регистр (по умолчанию не учитывается)
  --show-analyses       печатать морфологические разборы совпавших токенов
  --output FILE.xlsx    сохранить результаты в Excel вместо печати в консоль
  --output FILE.csv     сохранить результаты в CSV

Пример
  python cql_search.py corpus.jsonl '[lemma="пыкирык"]' --output results.xlsx
"""

import argparse
import csv
import json
import re
import sys
from typing import List, Optional, Tuple

# Загрузка корпуса

def load_corpus(path):
    with open(path, "r", encoding="utf-8") as f:
        text = f.read()

    decoder = json.JSONDecoder()
    records = []
    i = 0
    n = len(text)
    while i < n:
        while i < n and text[i] in " \t\r\n":
            i += 1
        if i >= n:
            break
        obj, end = decoder.raw_decode(text, i)
        records.append(obj)
        i = end
    return records


# Разбор морфологии

_TOKEN_RE = re.compile(r"<[^<>]+>|[^<>]+")

def parse_analysis(analysis: str) -> Optional[List[Tuple[str, List[str]]]]:
    if analysis.endswith("+?"):
        return None

    components: List[Tuple[str, List[str]]] = []
    for part in analysis.split("+"):
        if not part:
            continue
        pieces = _TOKEN_RE.findall(part)
        stems = [p for p in pieces if not p.startswith("<")]
        tags = [p[1:-1] for p in pieces if p.startswith("<")]
        stem = "".join(stems)
        if stem or tags:
            components.append((stem, tags))

    if not components:
        return None
    return components


def build_token(word, analyses, case_sensitive):
    norm_word = word if case_sensitive else word.lower()
    lemmas = set()
    poses = set()
    tags = set()
    has_analysis = False

    for analysis in analyses or []:
        comps = parse_analysis(analysis)
        if comps is None:
            continue
        has_analysis = True
        for stem, ctags in comps:
            if stem:
                lemmas.add(stem if case_sensitive else stem.lower())
            if ctags:
                poses.add(ctags[0])
                tags.update(ctags)

    if not has_analysis:
        # неразобранное слово — лемма считается равной словоформе
        lemmas.add(norm_word)

    return {
        "word": word,
        "word_norm": norm_word,
        "lemmas": lemmas,
        "pos": poses,
        "tags": tags,
        "analyses": analyses or [],
    }

# Парсер запроса

class QueryError(ValueError):
    pass

_QTOK_RE = re.compile(r"""
    \s*(?:
        (?P<lbrack> \[ ) |
        (?P<rbrack> \] ) |
        (?P<amp> & ) |
        (?P<pipe> \| ) |
        (?P<neq> != ) |
        (?P<eq> = ) |
        (?P<qstr> "(?:[^"\\]|\\.)*" ) |
        (?P<attr> [A-Za-z_]+ ) |
        (?P<quant> \{\s*\d+\s*(?:,\s*\d*)?\s*\} | [?*+] )
    )
""", re.VERBOSE)


def _unquote(s):
    return s[1:-1].replace('\\"', '"')


def tokenize_query(q):
    toks = []
    i = 0
    while i < len(q):
        if q[i].isspace():
            i += 1
            continue
        m = _QTOK_RE.match(q, i)
        if not m or m.end() == i:
            raise QueryError("Parsing error at: %r" % q[i:i + 20])
        kind = m.lastgroup
        val = m.group(kind)
        toks.append((kind, val))
        i = m.end()
    return toks


def parse_quant(qtext):
    if qtext == "?":
        return (0, 1)
    if qtext == "*":
        return (0, None)
    if qtext == "+":
        return (1, None)
    m = re.match(r"\{\s*(\d+)\s*(?:,\s*(\d*))?\s*\}", qtext)
    lo = int(m.group(1))
    if m.group(2) is None:
        hi = lo
    elif m.group(2) == "":
        hi = None
    else:
        hi = int(m.group(2))
    return (lo, hi)


def parse_query(q):
    toks = tokenize_query(q)
    i = 0
    pattern = []

    def peek():
        return toks[i] if i < len(toks) else (None, None)

    while i < len(toks):
        kind, val = toks[i]
        if kind == "lbrack":
            i += 1
            constraints = [[]]
            while True:
                kind, val = peek()
                if kind != "attr":
                    raise QueryError("Ожидалось имя атрибута (word/lemma/pos/tag) в позиции %d" % i)
                attr = val.lower()
                if attr not in ("word", "lemma", "pos", "tag"):
                    raise QueryError('Неизвестный атрибут "%s" (допустимы: word, lemma, pos, tag)' % attr)
                i += 1
                kind, val = peek()
                if kind not in ("eq", "neq"):
                    raise QueryError('Ожидался "=" или "!=" после атрибута "%s"' % attr)
                negate = kind == "neq"
                i += 1
                kind, val = peek()
                if kind != "qstr":
                    raise QueryError('Ожидалась строка в кавычках после %s=' % attr)
                pat = _unquote(val)
                i += 1
                constraints[-1].append((attr, pat, negate))
                kind, val = peek()
                if kind == "amp":
                    i += 1
                    continue
                if kind == "pipe":
                    i += 1
                    constraints.append([])
                    continue
                break
            kind, val = peek()
            if kind != "rbrack":
                raise QueryError("Ожидалась закрывающая ]")
            i += 1
        elif kind == "qstr":
            constraints = [[("word", _unquote(val), False)]]
            i += 1
        else:
            raise QueryError("Ожидался токен [attr=\"...\"] или \"строка\", получено: %r" % val)

        mn, mx = 1, 1
        kind, val = peek()
        if kind == "quant":
            mn, mx = parse_quant(val)
            i += 1
        pattern.append((constraints, mn, mx))

    if not pattern:
        raise QueryError("Пустой запрос")
    return pattern

# Матчинг запроса 

def compile_constraints(constraints, case_sensitive):
    flags = 0 if case_sensitive else re.IGNORECASE
    return [
        [(attr, re.compile(pat, flags), negate) for attr, pat, negate in alt]
        for alt in constraints
    ]


def _values(token, attr):
    if attr == "word":
        return (token["word_norm"],)
    if attr == "lemma":
        return token["lemmas"]
    if attr == "pos":
        return token["pos"]
    return token["tags"]


def _alt_matches(token, alt):
    for attr, rx, negate in alt:
        found = any(rx.fullmatch(v) for v in _values(token, attr))
        if found == negate:      
            return False
    return True


def token_matches(token, compiled_constraints):
    return any(_alt_matches(token, alt) for alt in compiled_constraints)


def compile_pattern(pattern, case_sensitive):
    return [(compile_constraints(c, case_sensitive), mn, mx) for c, mn, mx in pattern]


def find_matches(tokens, compiled_pattern):
    n = len(tokens)
    results = []
    start = 0
    while start <= n:
        best_end = None

        def rec(pi, ti):
            nonlocal best_end
            if pi == len(compiled_pattern):
                if best_end is None or ti > best_end:
                    best_end = ti
                return
            constraints, mn, mx = compiled_pattern[pi]
            t = ti
            c = 0
            if mn == 0:
                rec(pi + 1, ti)
            while (mx is None or c < mx) and t < n and token_matches(tokens[t], constraints):
                t += 1
                c += 1
                if c >= mn:
                    rec(pi + 1, t)

        rec(0, start)
        if best_end is not None and best_end > start:
            results.append((start, best_end))
            start = best_end
        else:
            start += 1
    return results


# Основная программа


def highlight(tokens, start, end):
    words = [t["word"] for t in tokens]
    words[start] = "**" + words[start]
    words[end - 1] = words[end - 1] + "**"
    return " ".join(words)


def collect_hits(records, compiled_pattern, case_sensitive):
    hits = []
    for rec_idx, rec in enumerate(records):
        morph = rec.get("morphology") or []
        if not morph:
            continue
        tokens = [
            build_token(m.get("form", ""), m.get("analyses"), case_sensitive)
            for m in morph
        ]
        matches = find_matches(tokens, compiled_pattern)
        for (start, end) in matches:
            matched_tokens = tokens[start:end]
            hits.append({
                "record": rec,
                "sent_index": rec_idx,
                "sentence_highlighted": highlight(tokens, start, end),
                "sentence": rec.get("ckt", ""),
                "match": " ".join(t["word"] for t in matched_tokens),
                "ru": rec.get("ru", "") or "",
                "source": rec.get("source", "") or "",
                "lemmas": "; ".join(sorted(set().union(*[t["lemmas"] for t in matched_tokens]))) if matched_tokens else "",
                "pos": "; ".join(sorted(set().union(*[t["pos"] for t in matched_tokens]))) if matched_tokens else "",
                "tags": "; ".join(sorted(set().union(*[t["tags"] for t in matched_tokens]))) if matched_tokens else "",
                "analyses": " | ".join(
                    a for t in matched_tokens for a in t["analyses"]
                ),
            })
    return hits


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def build_table(hits):
    keys = []
    for hit in hits:
        for k in hit["record"]:
            if k not in keys:
                keys.append(k)
    header = keys + ["match", "lemma"]
    rows = [
        [_cell(hit["record"].get(k)) for k in keys] + [hit["match"], hit["lemmas"]]
        for hit in hits
    ]
    return header, rows


def write_xlsx(hits, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Для сохранения в .xlsx нужен пакет openpyxl: pip install openpyxl",
              file=sys.stderr)
        sys.exit(1)

    header, rows = build_table(hits)
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)

    for col_idx, name in enumerate(header, start=1):
        width = 60 if name in ("ckt", "ru") else 30 if name in ("match", "lemma") else 18
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def write_csv(hits, path):
    header, rows = build_table(hits)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def main():
    ap = argparse.ArgumentParser(
        description="simplified CQL-search",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    ap.add_argument("corpus", help="путь к файлу корпуса (JSON/JSONL)")
    ap.add_argument("query", help='CQL-запрос, например \'[lemma="пыкирык"]\'')
    ap.add_argument("--limit", type=int, default=None, help="показать/сохранить не больше N совпадений")
    ap.add_argument("--case-sensitive", action="store_true", help="учитывать регистр")
    ap.add_argument("--show-analyses", action="store_true", help="печатать разборы совпавших токенов в консоль")
    ap.add_argument("--output", "-o", metavar="FILE", default=None,
                     help="сохранить результаты в файл вместо печати в консоль (.xlsx или .csv)")
    args = ap.parse_args()

    try:
        pattern = parse_query(args.query)
    except QueryError as e:
        print("Ошибка в запросе: %s" % e, file=sys.stderr)
        sys.exit(1)

    compiled_pattern = compile_pattern(pattern, args.case_sensitive)

    try:
        records = load_corpus(args.corpus)
    except FileNotFoundError:
        print("Файл не найден: %s" % args.corpus, file=sys.stderr)
        sys.exit(1)
    except json.JSONDecodeError as e:
        print("Ошибка чтения JSON: %s" % e, file=sys.stderr)
        sys.exit(1)

    hits = collect_hits(records, compiled_pattern, args.case_sensitive)
    total_hits = len(hits)
    shown_hits = hits if args.limit is None else hits[:args.limit]

    if args.output:
        if args.output.lower().endswith(".xlsx"):
            write_xlsx(shown_hits, args.output)
        elif args.output.lower().endswith(".csv"):
            write_csv(shown_hits, args.output)
        else:
            print("Неподдерживаемое расширение файла (выберите .xlsx или .csv): %s" % args.output,
                  file=sys.stderr)
            sys.exit(1)
        print("Сохранено %d из %d совпадений в %s" % (len(shown_hits), total_hits, args.output))
    else:
        for hit in shown_hits:
            print("-" * 70)
            print("[%d] %s" % (hit["sent_index"], hit["sentence_highlighted"]))
            if hit["ru"]:
                print("    RU: %s" % hit["ru"])
            if hit["source"]:
                print("    source: %s" % hit["source"])
            if args.show_analyses and hit["analyses"]:
                print("      %s" % hit["analyses"])
        print("=" * 70)
        print("Найдено совпадений: %d (показано: %d) из %d предложений" % (
            total_hits, len(shown_hits), len(records)))


if __name__ == "__main__":
    main()